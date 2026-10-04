"""Export only the requested public index's drawdown context, never the universe."""

import argparse
from datetime import date, datetime, timezone
import json
import math
from pathlib import Path
import tempfile


ROOT = Path(__file__).resolve().parent.parent
SOURCE_DIR = ROOT / "Portfolio_analysis/Output/Historical_drawdown_probability"
DESTINATION = ROOT / "HTML_portfolio/data/index_snapshot.json"
SYMBOL = "CSSPX.MI"


def build_snapshot(rows, expected_date=None):
    matches = [row for row in rows if row.get("ticker") == SYMBOL]
    if len(matches) != 1:
        raise ValueError("Expected exactly one CSSPX.MI observation")
    row = matches[0]
    score = row.get("invquant_dd")
    if isinstance(score, bool) or not isinstance(score, (int, float)):
        raise ValueError("Missing or nonnumeric index score")
    if not math.isfinite(score) or not 0 <= score <= 100:
        raise ValueError("Index score must be finite and between 0 and 100")
    raw_date = row.get("Reference_date")
    if isinstance(raw_date, datetime):
        raw_date = raw_date.date()
    analysis_date = raw_date.isoformat() if isinstance(raw_date, date) else raw_date
    # Reference_date is the analysis request date; do not label it a quote date.
    if not isinstance(analysis_date, str):
        raise ValueError("Missing analysis date")
    if date.fromisoformat(analysis_date).isoformat() != analysis_date:
        raise ValueError("Analysis date must use YYYY-MM-DD")
    if expected_date is not None and analysis_date != expected_date:
        raise ValueError("Analysis date disagrees with source filename")
    return {
        "schema_version": 1,
        "symbol": SYMBOL,
        "invquant_dd": float(score),
        "analysis_date": analysis_date,
        "generated_at": datetime.now(timezone.utc).isoformat(),
    }


def export_snapshot(rows, destination=DESTINATION, expected_date=None):
    payload = build_snapshot(rows, expected_date)
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=destination.parent,
                                         prefix=".index-snapshot-", suffix=".tmp", delete=False) as handle:
            temporary = Path(handle.name)
            json.dump(payload, handle, indent=2, allow_nan=False)
            handle.write("\n")
        temporary.replace(destination)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return payload


def export_latest(source_dir=SOURCE_DIR, destination=DESTINATION):
    from openpyxl import load_workbook

    dated_sources = []
    for path in Path(source_dir).glob("invquantiledd_*.xlsx"):
        stamp = path.stem[len("invquantiledd_"):]
        try:
            parsed = date.fromisoformat(stamp)
        except ValueError:
            continue
        if parsed.isoformat() == stamp:
            dated_sources.append((stamp, path))
    if not dated_sources:
        raise ValueError("No dated drawdown workbook found")
    stamp, path = max(dated_sources)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        values = workbook.active.iter_rows(values_only=True)
        headers = next(values)
        rows = [dict(zip(headers, row)) for row in values]
        return export_snapshot(rows, destination, stamp)
    finally:
        workbook.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, default=SOURCE_DIR)
    parser.add_argument("--output", type=Path, default=DESTINATION)
    args = parser.parse_args()
    result = export_latest(args.source_dir, args.output)
    print(f"Exported {result['symbol']} index snapshot ({result['analysis_date']})")
